import csv
import base64
import openpyxl
from io import StringIO, BytesIO
from asyncio import base_events
from django.conf import settings
from saleor.external_services.mail.constants import TEMPLATE_CODE
from saleor.external_services.mail.mail_manager import MailManager
from saleor.utilities.api_client import ApiClient
import logging

logger = logging.getLogger(__name__)

class MailImpl(MailManager):

    def _set_mail_header():
        header = {'Authorization': 'Bearer ' + settings.SENDGRID_API_KEY,
                  'Content-Type': 'application/json'}
        return header

    def _set_mail_body(mail_type, template_data, recipient_email, cc=None):

        if not cc:
            cc=[]

        if settings.IS_BETA :
            recipient_email = settings.BETA_RECIPIENT_EMAIL
        else:
            if not TEMPLATE_CODE[mail_type]['recipient_email']:
                recipient_email = recipient_email
            else:
                recipient_email = TEMPLATE_CODE[mail_type]['recipient_email']

        cc_list = []
        for email in cc:
            cc_list.append({"email": email})

        body = {
            "personalizations": [{
                                    "to": [{
                                        "email": recipient_email
                                    }],
                                    "dynamic_template_data": template_data                                  
                                }],
            "from": {
                "email": TEMPLATE_CODE[mail_type]['sender_email'], 
                "name": TEMPLATE_CODE[mail_type]['sender_name']
            },
            "reply_to": {
                "email": TEMPLATE_CODE[mail_type]['sender_email'], 
                "name": TEMPLATE_CODE[mail_type]['sender_name']
            },
            "template_id": TEMPLATE_CODE[mail_type]['template_id']
        }

        if cc_list:
            body["personalizations"][0].update({'cc': cc_list})

        return body

    def _set_mail_body_without_template(mail_type, mail_text, subject, recipient_email, cc=None):

        if not cc:
            cc=[]

        if settings.IS_BETA :
            recipient_email = settings.BETA_RECIPIENT_EMAIL
        else:
    
            if not TEMPLATE_CODE[mail_type]['recipient_email']:
                recipient_email = recipient_email
            else:
                recipient_email = TEMPLATE_CODE[mail_type]['recipient_email']

        cc_list = []
        for email in cc:
            cc_list.append({"email": email})

        body = {
            "from": {
                "email": TEMPLATE_CODE[mail_type]['sender_email'], 
                "name": TEMPLATE_CODE[mail_type]['sender_name']
            },
            "reply_to": {
                "email": TEMPLATE_CODE[mail_type]['sender_email'], 
                "name": TEMPLATE_CODE[mail_type]['sender_name']
            },
            "content": [{
                "type": "text/plain",
                "value": mail_text
            }],
            "personalizations": [
                {
                "to": [
                    {
                        "email": recipient_email
                    }
                ],
                "subject": subject
                }
            ]
        }

        if cc_list:
            body["personalizations"][0].update({'cc': cc_list})
        
        return body


    def _set_mail_body_with_attachment(mail_type, mail_text, subject, recipient_email, attachment_list, cc=None):

        if not cc:
            cc=[]

        if settings.IS_BETA :
            recipient_email = settings.BETA_RECIPIENT_EMAIL
        else:
    
            if not TEMPLATE_CODE[mail_type]['recipient_email']:
                recipient_email = recipient_email
            else:
                recipient_email = TEMPLATE_CODE[mail_type]['recipient_email']

        cc_list = []
        for email in cc:
            cc_list.append({"email": email})

        body = {
            "from": {
                "email": TEMPLATE_CODE[mail_type]['sender_email'], 
                "name": TEMPLATE_CODE[mail_type]['sender_name']
            },
            "reply_to": {
                "email": TEMPLATE_CODE[mail_type]['sender_email'], 
                "name": TEMPLATE_CODE[mail_type]['sender_name']
            },
            "content": [{
                "type": "text/plain",
                "value": mail_text
            }],
            "personalizations": [
                {
                "to": [
                    {
                        "email": recipient_email
                    }
                ],
                "subject": subject
                }
            ],
            "attachments": attachment_list
        }

        if cc_list:
            body["personalizations"][0].update({'cc': cc_list})
        
        return body

    @classmethod
    def send_mail(cls, mail_type, template_data: dict, recipient_email='', cc = None):
        """
        sends mail using templates
        """

        if not cc:
            cc=[]

        mail_url = settings.SENDGRID_API_URL
        header = cls._set_mail_header()
        body = cls._set_mail_body(mail_type, template_data, recipient_email, cc=cc)
        recipient_email = body['personalizations'][0]['to'][0]['email']

        logger.info(
            "Send Mail with template for the recipient :: %s having header :: %s and body :: %s ", recipient_email, header, body
        )
        api_client = ApiClient(url = mail_url)
        api_client.update_headers(header)
        api_client.update_body(body)
        api_client.post()
        logger.info(
            "ApiResponse for Send Mail with template for the recipient :: %s is :: %s ",
            recipient_email, api_client.fetch_response_code()
        )
        return api_client.fetch_response_code()

    @classmethod
    def send_mail_without_template(cls, mail_type, mail_text, subject, recipient_email='', cc = None):

        if not cc:
            cc=[]
            
        mail_url = settings.SENDGRID_API_URL
        header = cls._set_mail_header()
        body = cls._set_mail_body_without_template(mail_type, mail_text, subject, recipient_email, cc=cc)
        api_client = ApiClient(url = mail_url)
        api_client.update_headers(header)
        api_client.update_body(body)
        api_client.post()
        print(api_client.fetch_response_code())
        return api_client.fetch_response_code()

    @classmethod
    def send_mail_with_attachment(cls, mail_type, mail_text, subject, attachment_list, recipient_email='', cc = None):
        '''
        attachment_list structure ::

            "attachments": [
                {"content": "BASE64_ENCODED_CONTENT", "type": "application/csv", "filename": "attachment.txt"},
                {"content": "BASE64_ENCODED_CONTENT", "type": "application/csv", "filename": "attachment.txt"},
                ...
                ]
                
        '''
        if not cc:
            cc=[]
            
        mail_url = settings.SENDGRID_API_URL
        header = cls._set_mail_header()
        body = cls._set_mail_body_with_attachment(mail_type, mail_text, subject, recipient_email, attachment_list, cc=cc)
        api_client = ApiClient(url = mail_url)
        api_client.update_headers(header)
        api_client.update_body(body)
        api_client.post()
        print(api_client.fetch_response_code())
        return api_client.fetch_response_code()

    @classmethod
    def get_csv_attachment_dict(self, filename, rows, fieldnames=None):
        f = StringIO()
        if fieldnames:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
        else:
            csv.writer(f).writerows(rows)
        base_encoded_file = base64.b64encode(f.getvalue().encode()).decode()
        return {"content": base_encoded_file, "type": "application/csv", "filename": filename}

    @classmethod
    def get_excel_attachment_dict(self, filename, sheets:dict):
        """
        sheets = { sheet_name: rows }
        """
        f = BytesIO()
        workbook = openpyxl.Workbook()

        for sheet_name, rows in sheets.items():
            workbook.create_sheet(sheet_name)
            sheet = workbook[sheet_name]
            for i, row in enumerate(rows, start=1):
                for j, val in enumerate(row, start=1):
                    sheet.cell(row=i, column=j).value = val
        
        workbook.remove(workbook['Sheet'])
        workbook.save(f)

        base_encoded_file = base64.b64encode(f.getvalue()).decode()
        return {"content": base_encoded_file, "type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "filename": filename}

